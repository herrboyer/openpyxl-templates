from collections import Counter, namedtuple
from collections import deque
from enum import Enum
from itertools import chain, zip_longest, repeat

from clint import OrderedDict
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

from openpyxl_templates.exceptions import OpenpyxlTemplateCellException, CellExceptions, \
    RowExceptions, SheetException
from openpyxl_templates.table_sheet.columns import TableColumn
from openpyxl_templates.templated_sheet import TemplatedSheet
from openpyxl_templates.utils import Typed, MAX_COLUMN_INDEX


class TableSheetException(SheetException):
    pass


class ColumnHeadersNotUnique(TableSheetException):
    def __init__(self, columns):
        counter = Counter(column.header for column in columns)
        super().__init__("headers '%s' has been declared more then once in the same TableSheet" % tuple(
            header
            for (header, count)
            in counter.items()
            if count > 1
        ))


class TempleteStyleNotFound(TableSheetException):
    def __init__(self, missing_style_name, style_set):
        super().__init__(
            "The style '%s' has not been declared. Avaliable styles are: %s)"
            % (missing_style_name, style_set.names)
        )


class NoTableColumns(TableSheetException):
    def __init__(self, table_sheet):
        super().__init__(
            "The TableSheet '%s' has no columns. Declare atleast one."
            % table_sheet.sheetname
        )


class HeadersNotFound(TableSheetException):
    def __init__(self, table_sheet):
        super().__init__(
            "Header column not found on sheet '%s' either make sure that the following headers are "
            "present '%s'." % (
                table_sheet.sheetname,
                ", ".join(table_sheet.headers)
            )
        )


class TableSheetExceptionPolicy(Enum):
    RaiseCellException = 1
    RaiseRowException = 2
    RaiseSheetException = 3
    IgnoreRow = 4


class TableSheet(TemplatedSheet):
    item_class = TableColumn

    title_style = Typed("title_style", expected_type=str, value="Title")
    description_style = Typed("description_style", expected_type=str, value="Description")

    format_as_table = Typed("format_as_header", expected_type=bool, value=True)
    freeze_header = Typed("freeze_header", expected_type=bool, value=True)
    hide_excess_columns = Typed("hide_excess_columns", expected_type=bool, value=True)

    _first_data_cell = None
    _last_data_cell = None
    _first_header_cell = None
    _last_header_cell = None

    def __init__(self, sheetname=None, active=None):
        super().__init__(sheetname=sheetname, active=active)

        self.columns = []
        index = 1
        for object_attribute, column in self._items.items():
            if not column._object_attribute:
                column._object_attribute = object_attribute
            column.column_index = index
            index += 1

            self.columns.append(column)

        self._row_class = namedtuple(
            "%sRow" % self.__class__.__name__,
            (column.object_attribute for column in self.columns)
        )

        self._validate()

    def _validate(self):
        self._check_atleast_one_column()
        self._check_unique_column_headers()

    def _check_atleast_one_column(self):
        if not self.columns:
            raise NoTableColumns(self)

    def _check_unique_column_headers(self):
        if len(set(column.header for column in self.columns)) < len(self.columns):
            raise ColumnHeadersNotUnique(self.columns)

    def write(self, title=None, description=None, objects=None):
        worksheet = self.worksheet

        self.prepare_worksheet(worksheet)
        self.write_title(worksheet, title)
        self.write_description(worksheet, description)
        self.write_headers(worksheet)
        self.write_rows(worksheet, objects)
        self.post_process_worksheet(worksheet)

    def prepare_worksheet(self, worksheet):
        for column in self.columns:
            column.prepare_worksheet(worksheet)

        # Register styles
        style_names = set(chain(
            (self.title_style, self.description_style),
            *((column.row_style, column.header_style) for column in self.columns)
        ))

        existing_names = set(self.workbook.named_styles)

        for name in style_names:
            if name in existing_names:
                continue

            if name not in self.workbook.template_styles:
                raise TempleteStyleNotFound(name, self.workbook.template_styles)

            self.workbook.add_named_style(self.workbook.template_styles[name])

    def write_title(self, worksheet, title=None):
        if not title:
            return

        title = WriteOnlyCell(ws=worksheet, value=title)
        title.style = self.title_style

        worksheet.append((title,))

    def write_description(self, worksheet, description=None):
        if not description:
            return

        description = WriteOnlyCell(ws=worksheet, value=description)
        description.style = self.description_style

        worksheet.append((description,))

    def write_headers(self, worksheet):
        headers = tuple(
            column.create_header(worksheet)
            for column in self.columns
        )

        self.worksheet.append(headers)

        self._first_header_cell = headers[0]
        self._last_header_cell = headers[-1]

    def write_rows(self, worksheet, objects=None):
        self._first_data_cell = None
        cells = None
        for obj in objects:
            cells = tuple(column.create_cell(worksheet, column.get_value_from_object(obj)) for column in self.columns)
            worksheet.append(cells)

            if not self._first_data_cell:
                self._first_data_cell = cells[0]

            for cell, column in zip(cells, self.columns):
                column.post_process_cell(worksheet, cell)

        if cells:
            self._last_data_cell = cells[-1]

    def post_process_worksheet(self, worksheet):
        for column in self.columns:
            column.post_process_worksheet(worksheet)

        if self.active:
            self.activate()

        if self.format_as_table:
            worksheet.add_table(
                Table(
                    ref="%s:%s" % (
                        self._first_header_cell.coordinate,
                        self._last_data_cell.coordinate if self._last_data_cell else self._last_header_cell.coordinate
                    ),
                    displayName=self.sheetname,
                )
            )

        if self.freeze_header:
            worksheet.freeze_panes = self._first_data_cell or self._first_header_cell

        if self.hide_excess_columns:
            for i in range(len(self.columns) + 1, MAX_COLUMN_INDEX + 1):
                worksheet.column_dimensions[get_column_letter(i)].hidden = True

    def read(self, exception_policy=TableSheetExceptionPolicy.RaiseCellException, look_for_header=True):
        header_found = not look_for_header

        rows = self.worksheet.__iter__()
        try:
            while not header_found:
                header_found = self._is_row_header(rows.__next__())

            row_exceptions = []
            while True:
                try:
                    yield self.object_from_row(rows.__next__(), exception_policy=exception_policy)
                except CellExceptions as e:
                    if exception_policy <= TableSheetExceptionPolicy.RaiseRowException:
                        raise e
                    else:
                        row_exceptions.append(e)

                if row_exceptions and exception_policy <= TableSheetExceptionPolicy.RaiseSheetException:
                    raise RowExceptions(row_exceptions)
        except StopIteration:
            pass

        if not header_found:
            raise HeadersNotFound(self)

    def _is_row_header(self, row):
        for cell, header in zip(chain(row, repeat(None)), self.headers):
            if str(cell.value) != header:
                return False
        return True

    def object_from_row(self, row, exception_policy=TableSheetExceptionPolicy.RaiseCellException):
        data = OrderedDict()
        cell_exceptions = []
        for cell, column in zip(chain(row, repeat(None)), self.columns):
            try:
                data[column.object_attribute] = column.from_excel_with_blank_check(cell)
            except OpenpyxlTemplateCellException as e:
                if exception_policy <= TableSheetExceptionPolicy.RaiseCellException:
                    raise e
                else:
                    cell_exceptions.append(e)

        if cell_exceptions:
            raise CellExceptions(cell_exceptions)

        return self.create_object(data)

    def create_object(self, data):
        return self._row_class(*data.values())

    @property
    def headers(self):
        return (column.header for column in self.columns)
